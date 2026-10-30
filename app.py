import os
from io import BytesIO
import streamlit as st
import pandas as pd
from datetime import datetime
import json

# CSS code to hide Streamlit branding
hide_streamlit_style = """
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    </style>
"""
st.markdown(hide_streamlit_style, unsafe_allow_html=True)

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image as RLImage
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from docx import Document

st.set_page_config(page_title="Residence Verification Report Generator", layout="centered")

st.title("🏡 Residence Verification Report Generator")

# ----------------- PERMANENT JSON STORAGE FOR HISTORY -----------------
HISTORY_FILE = "report_history.json"

def load_history():
    if os.path.exists(HISTORY_FILE):
        try:
            with open(HISTORY_FILE, "r", encoding="utf-8") as f:
                content = f.read()
                if content.strip():
                    return json.loads(content)
        except Exception as e:
            return []
    else:
        try:
            with open(HISTORY_FILE, "w", encoding="utf-8") as f:
                json.dump([], f)
        except:
            pass
    return []

def save_history(history_list):
    serializable_list = []
    for item in history_list:
        serializable_list.append({
            "Timestamp": item["Timestamp"],
            "Applicant Name": item["Applicant Name"],
            "Branch Name": item["Branch Name"],
            "Customer Mobile": item["Customer Mobile"],
            "FullData": item["FullData"]
        })
    try:
        with open(HISTORY_FILE, "w", encoding="utf-8") as f:
            json.dump(serializable_list, f, ensure_ascii=False, indent=4)
    except Exception as e:
        pass

if 'report_history' not in st.session_state:
    st.session_state.report_history = load_history()

if 'is_authenticated' not in st.session_state:
    st.session_state.is_authenticated = False

total_reports = len(st.session_state.report_history)
st.info(f"📊 Muth bonuwa report-or songkha (Total Reports Generated): **{total_reports}**")

st.subheader("⚙️ Quick Status Selection")
col_s1, col_s2 = st.columns(2)
with col_s1:
    solar_install = st.selectbox("Solar Installed?", ["Yes", "No"], key="solar_select")
with col_s2:
    person_met = st.text_input("Person Met *", "OSMAN ALI", key="person_met_quick")

if solar_install == "No":
    auto_verifier_remarks = f"I Visited At The Provided Address & I Met with {person_met} Confirmed His Address. Neighbor’s feedback Is Satisfactory. There is no Adverse report found as per enquiry. Hence The Residence Report is Positive."
else:
    auto_verifier_remarks = "I Visited At The Provided Address and found solar panel is successfully installed"

with st.form("verification_form"):
    st.subheader("📋 Applicant & Visit Details (Fields marked with * are mandatory)")
    col1, col2 = st.columns(2)
    
    with col1:
        branch_name = st.text_input("Branch Name *", "SBI SAMAGURI BRANCH")
        lifting_date = st.text_input("Lifting Date *", "03.10.2026")
        applicant_name = st.text_input("Name of Applicant *", "USUF ALI")
        customer_mobile = st.text_input("Customer Mobile Number *", "9876543210")
        occupation = st.text_input("Occupation", "BUSINESS")
        address = st.text_area("Residence Address *", "VILL:LALUNG GAON, P/O-KASORI P/S-JURIA PINCODE-782124 DIST: NAGAON, ASSAM")
        date_of_visit = st.text_input("Date of Visit *", "03.10.2026")
        time_of_visit = st.text_input("Time of Visit *", "09:33:00 AM")
        address_confirmed = st.text_input("Address Confirmed", "YES")
        relationship = st.text_input("Relationship *", "FATHER")
        marital_status = st.text_input("Marital Status", "MARRIED")
        family_members = st.text_input("Total Family Members", "05")
        earning_members = st.text_input("Total Earning Members", "03")
        children = st.text_input("Children", "00")
        
    with col2:
        is_spouse_working = st.text_input("Is Spouse Working? (Y/N)", "NA")
        spouse_details = st.text_input("If Yes Then Details", "NA")
        ownership = st.text_input("Ownership", "OWN HOUSE")
        years_in_house = st.text_input("No of Years in Existing House", "BY BIRTH")
        locality = st.text_input("Locality", "NA")
        unit = st.text_input("Unit", "GROUND FLOOR")
        construction = st.text_input("Construction", "CEMENTED")
        roof_type = st.text_input("Type of Roof", "CONCRETE")
        approx_area = st.text_area("Approx Area of the House (Sq Ft.)", "1000 SQFT")
        floor = st.text_input("Floor of the House", "GROUND")
        assets_seen = st.text_input("Assets Seen? (Y/N)", "YES")
        vehicle = st.text_input("Vehicle", "NA")
        portrait = st.text_input("Portrait of Political Leaders", "NO")

    st.subheader("📊 Additional Details & Remarks")
    col3, col4 = st.columns(2)
    with col3:
        purpose = st.text_input("Purpose", "SURYA GHAR LOAN")
        prop_location = st.text_input("Proposed Property Location", "EASY")
        loans_taken = st.text_input("Any Loans Taken? (Y/N)", "NO")
        financier = st.text_input("Financier", "NA")
        emi = st.text_input("EMI", "NA")
        accessibility = st.text_input("Accessibility of Office", "Easy")
    with col4:
        neighbor_name = st.text_input("Neighbor Name", "RAHIM UDDIN")
        neighbor_mobile = st.text_input("Neighbor Mobile Number", "9123456789")
        neighbors_feedback = st.text_input("Neighbor's Feedback", "GOOD")
        
        verifier_remarks = st.text_input("Verifier Remarks", value=auto_verifier_remarks)
        
        remarks = st.text_input("Remarks", "Positive")
        prev_applied = st.text_input("Did Customer Apply Previously?", "NA")
        prev_date = st.text_input("Date of Previous File", "NA")
        recommendation = st.text_input("Recommendation", "Yes")
        verifier_name = st.text_input("Verifier Name", "EKRAMUL HUSSAIN")

    st.subheader("📸 Upload Verification Photos (Dukhan photo upload korok)")
    uploaded_photo1 = st.file_uploader("Upload House/Location Photo 1", type=["jpg", "jpeg", "png"], key="p1")
    uploaded_photo2 = st.file_uploader("Upload Customer/Verifier Photo 2", type=["jpg", "jpeg", "png"], key="p2")

    st.subheader("🔵 Upload Authorized Seal Image")
    uploaded_seal = st.file_uploader("Upload Seal Image (PNG/JPG)", type=["jpg", "jpeg", "png"], key="seal")

    submitted = st.form_submit_button("Generate Files")

def generate_report_files(data_dict, p1, p2, seal):
    pdf_buffer = BytesIO()
    doc = SimpleDocTemplate(pdf_buffer, pagesize=letter, rightMargin=20, leftMargin=20, topMargin=15, bottomMargin=15)
    story = []
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('TitleStyle', parent=styles['Heading1'], fontSize=12, textColor=colors.HexColor('#006633'), alignment=1, spaceAfter=2)
    subtitle_style = ParagraphStyle('SubTitleStyle', parent=styles['Normal'], fontSize=7.5, textColor=colors.HexColor('#555555'), alignment=1, spaceAfter=6)
    cell_style = ParagraphStyle('CellStyle', parent=styles['Normal'], fontSize=8, leading=10)
    
    story.append(Paragraph("<b>TIGER 4 INDIA LIMITED</b>", title_style))
    story.append(Paragraph("Corporate Office: #7, Lower Ground Floor, L.S.C., B-1, Vasant Kunj, New Delhi-110070 India<br/>E-mail id: tigindialtd@gmail.com | Mobile: 8282864451", subtitle_style))
    story.append(Paragraph("<b><u>RESIDENCE VERIFICATION REPORT</u></b>", ParagraphStyle('H2', parent=title_style, fontSize=10, textColor=colors.HexColor('#222222'))))
    story.append(Spacer(1, 4))
    
    table_data = [
        [Paragraph("<b>Branch Name:</b>", cell_style), Paragraph(data_dict['branch_name'], cell_style), Paragraph("<b>Lifting Date:</b>", cell_style), Paragraph(data_dict['lifting_date'], cell_style)],
        [Paragraph("<b>Applicant Name:</b>", cell_style), Paragraph(data_dict['applicant_name'], cell_style), Paragraph("<b>Customer Mobile:</b>", cell_style), Paragraph(data_dict['customer_mobile'], cell_style)],
        [Paragraph("<b>Occupation:</b>", cell_style), Paragraph(data_dict['occupation'], cell_style), Paragraph("<b>Date of Visit:</b>", cell_style), Paragraph(data_dict['date_of_visit'], cell_style)],
        [Paragraph("<b>Address:</b>", cell_style), Paragraph(data_dict['address'], cell_style), Paragraph("<b>Time of Visit:</b>", cell_style), Paragraph(data_dict['time_of_visit'], cell_style)],
        [Paragraph("<b>Address Confirmed:</b>", cell_style), Paragraph(data_dict['address_confirmed'], cell_style), Paragraph("<b>Person Met:</b>", cell_style), Paragraph(data_dict['person_met'], cell_style)],
        [Paragraph("<b>Relationship:</b>", cell_style), Paragraph(data_dict['relationship'], cell_style), Paragraph("<b>Marital Status:</b>", cell_style), Paragraph(data_dict['marital_status'], cell_style)],
        [Paragraph("<b>Family / Earning:</b>", cell_style), Paragraph(f"Family: {data_dict['family_members']} | Earning: {data_dict['earning_members']}", cell_style), Paragraph("<b>Children:</b>", cell_style), Paragraph(data_dict['children'], cell_style)],
        [Paragraph("<b>Spouse Working:</b>", cell_style), Paragraph(data_dict['is_spouse_working'], cell_style), Paragraph("<b>Details:</b>", cell_style), Paragraph(data_dict['spouse_details'], cell_style)],
        [Paragraph("<b>Ownership:</b>", cell_style), Paragraph(data_dict['ownership'], cell_style), Paragraph("<b>Years in House:</b>", cell_style), Paragraph(data_dict['years_in_house'], cell_style)],
        [Paragraph("<b>Construction:</b>", cell_style), Paragraph(data_dict['construction'], cell_style), Paragraph("<b>Roof Type:</b>", cell_style), Paragraph(data_dict['roof_type'], cell_style)],
        [Paragraph("<b>Approx Area:</b>", cell_style), Paragraph(data_dict['approx_area'], cell_style), Paragraph("<b>Floor:</b>", cell_style), Paragraph(data_dict['floor'], cell_style)],
        [Paragraph("<b>Assets Seen:</b>", cell_style), Paragraph(data_dict['assets_seen'], cell_style), Paragraph("<b>Vehicle:</b>", cell_style), Paragraph(data_dict['vehicle'], cell_style)],
        [Paragraph("<b>Political Portrait:</b>", cell_style), Paragraph(data_dict['portrait'], cell_style), Paragraph("<b>Purpose:</b>", cell_style), Paragraph(data_dict['purpose'], cell_style)],
        [Paragraph("<b>Loans Taken:</b>", cell_style), Paragraph(data_dict['loans_taken'], cell_style), Paragraph("<b>Financier / EMI:</b>", cell_style), Paragraph(f"{data_dict['financier']} / {data_dict['emi']}", cell_style)],
        [Paragraph("<b>Neighbor Name:</b>", cell_style), Paragraph(data_dict['neighbor_name'], cell_style), Paragraph("<b>Neighbor Mobile:</b>", cell_style), Paragraph(data_dict['neighbor_mobile'], cell_style)],
        [Paragraph("<b>Neighbor Feedback:</b>", cell_style), Paragraph(data_dict['neighbors_feedback'], cell_style), Paragraph("<b>Solar Installed:</b>", cell_style), Paragraph(data_dict['solar_install'], cell_style)],
        [Paragraph("<b>Remarks:</b>", cell_style), Paragraph(f"<b><font color='green'>{data_dict['remarks']}</font></b>", cell_style), Paragraph("<b>Status:</b>", cell_style), Paragraph("Positive", cell_style)]
    ]
    
    t = Table(table_data, colWidths=[110, 160, 100, 202])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F4F9F4')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#006633')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#C2D6C2')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2),
    ]))
    
    story.append(t)
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>Verifier Remarks:</b> {data_dict['final_verifier_remarks']}", cell_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("<b>Verification Photos:</b>", ParagraphStyle('H3', parent=styles['Heading3'], fontSize=9, textColor=colors.HexColor('#006633'))))
    story.append(Spacer(1, 2))
    
    img_data = []
    temp_files = []
    for idx, up_file in enumerate([p1, p2]):
        if up_file is not None:
            temp_path = f"temp_img_{idx}_{datetime.now().microsecond}.jpg"
            with open(temp_path, "wb") as f:
                f.write(up_file.getbuffer())
            temp_files.append(temp_path)
            img_data.append(RLImage(temp_path, width=220, height=130))
            
    if len(img_data) > 0:
        while len(img_data) < 2:
            img_data.append("")
        photo_table = Table([[img_data[0], img_data[1]]], colWidths=[286, 286])
        photo_table.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'CENTER'), ('VALIGN', (0,0), (-1,-1), 'MIDDLE')]))
        story.append(photo_table)

    story.append(Spacer(1, 6))
    
    if seal is not None:
        seal_path = f"temp_seal_{datetime.now().microsecond}.png"
        with open(seal_path, "wb") as f:
            f.write(seal.getbuffer())
        temp_files.append(seal_path)
        seal_img = RLImage(seal_path, width=80, height=80)
        name_para = Paragraph(f"<b>Verifier Name: - {data_dict['verifier_name']}</b><br/><br/><b>Authorized Signature with Seal</b>", cell_style)
        sig_table = Table([[seal_img], [name_para]], colWidths=[250])
        sig_table.setStyle(TableStyle([('ALIGN', (0,0), (-1,-1), 'LEFT'), ('VALIGN', (0,0), (-1,-1), 'MIDDLE'), ('BOTTOMPADDING', (0,0), (-1,0), -38)]))
    else:
        name_para = Paragraph(f"<b>Verifier Name: - {data_dict['verifier_name']}</b><br/><br/><b>Authorized Signature with Seal</b>", cell_style)
        sig_table = Table([[name_para]], colWidths=[250])

    story.append(sig_table)
    doc.build(story)
    pdf_buffer.seek(0)
    
    doc_word = Document()
    doc_word.add_heading("TIGER 4 INDIA LIMITED", level=1)
    doc_word.add_paragraph("Corporate Office: #7, Lower Ground Floor, L.S.C., B-1, Vasant Kunj, New Delhi-110070 India\nE-mail id: tigindialtd@gmail.com | Mobile: 8282864451")
    doc_word.add_heading("RESIDENCE VERIFICATION REPORT", level=2)
    doc_word.add_paragraph(f"Branch Name: {data_dict['branch_name']} | Lifting Date: {data_dict['lifting_date']}")
    doc_word.add_paragraph(f"Name of Applicant: {data_dict['applicant_name']} | Customer Mobile: {data_dict['customer_mobile']}")
    doc_word.add_paragraph(f"Occupation: {data_dict['occupation']}")
    doc_word.add_paragraph(f"Residence Address: {data_dict['address']}")
    doc_word.add_paragraph(f"Date of Visit: {data_dict['date_of_visit']} | Time: {data_dict['time_of_visit']}")
    doc_word.add_paragraph(f"Person Met: {data_dict['person_met']} ({data_dict['relationship']})")
    doc_word.add_paragraph(f"Family Members: {data_dict['family_members']} | Earning: {data_dict['earning_members']}")
    doc_word.add_paragraph(f"Ownership: {data_dict['ownership']} | Construction: {data_dict['construction']} | Roof: {data_dict['roof_type']}")
    doc_word.add_paragraph(f"Neighbor Name: {data_dict['neighbor_name']} | Neighbor Mobile: {data_dict['neighbor_mobile']}")
    doc_word.add_paragraph(f"Solar Installed: {data_dict['solar_install']}")
    doc_word.add_paragraph(f"Verifier Remarks: {data_dict['final_verifier_remarks']}")
    doc_word.add_paragraph(f"\nVerifier Name: - {data_dict['verifier_name']}\nAuthorized Signature with Seal")
    
    word_buffer = BytesIO()
    doc_word.save(word_buffer)
    word_buffer.seek(0)
    
    for tf in temp_files:
        if os.path.exists(tf):
            os.remove(tf)
            
    return pdf_buffer, word_buffer

if submitted:
    missing_fields = []
    if not branch_name.strip():
        missing_fields.append("Branch Name")
    if not lifting_date.strip():
        missing_fields.append("Lifting Date")
    if not applicant_name.strip():
        missing_fields.append("Name of Applicant")
    if not customer_mobile.strip():
        missing_fields.append("Customer Mobile Number")
    if not address.strip():
        missing_fields.append("Residence Address")
    if not date_of_visit.strip():
        missing_fields.append("Date of Visit")
    if not time_of_visit.strip():
        missing_fields.append("Time of Visit")
    if not relationship.strip():
        missing_fields.append("Relationship")
    if not person_met.strip():
        missing_fields.append("Person Met")

    if missing_fields:
        st.error(f"❌ Please fill out all mandatory fields marked with (*): {', '.join(missing_fields)}")
    else:
        if solar_install == "No":
